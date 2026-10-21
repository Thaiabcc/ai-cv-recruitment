import json
import base64
import unittest
from io import BytesIO
from unittest.mock import patch

import httpx
from docx import Document
from pydantic import SecretStr
from pypdf import PdfReader
from sqlalchemy import func, select
from sqlalchemy.orm import Session

import test_accounts
from app.config import settings
from app.models import CV, CVPhoto
from PIL import Image
from zipfile import ZipFile


class CVTests(unittest.TestCase):
    setUp = test_accounts.AccountTests.setUp
    tearDown = test_accounts.AccountTests.tearDown
    register = test_accounts.AccountTests.register

    def profile(self):
        self.register()
        payload = {'full_name': 'Nguyễn Minh Anh', 'phone': '0901234567', 'address': 'Hà Nội',
                   'education': 'Đại học Công nghệ thông tin', 'skills': 'Python, MySQL',
                   'experience': 'Thực tập phát triển backend', 'projects': 'Hệ thống quản lý CV',
                   'desired_position': 'Backend Intern'}
        self.assertEqual(self.client.put('/api/profile', json=payload).status_code, 200)
        return payload

    def draft(self):
        response = self.client.post('/api/cvs/draft')
        self.assertEqual(response.status_code, 201, response.text)
        return response.json()

    def career(self):
        return {'title': 'Backend Intern', 'summary': 'Ứng viên có kỹ năng Python và MySQL.',
                'education': 'Đại học Công nghệ thông tin', 'skills': 'Python, MySQL',
                'experience': 'Thực tập phát triển backend', 'projects': 'Hệ thống quản lý CV'}

    def ai_response(self, data=None, status=200, finish='STOP'):
        body = {'candidates': [{'finishReason': finish, 'content': {'parts': [
            {'text': json.dumps(data if data is not None else self.career(), ensure_ascii=False)}]}}]}
        return httpx.Response(status, json=body)

    def count(self):
        with Session(self.connection) as db:
            return db.scalar(select(func.count()).select_from(CV))

    def test_authentication_and_missing_profile(self):
        for path in ['/api/cvs', '/api/cvs/status', '/api/cvs/1', '/api/cvs/1/export/pdf']:
            self.assertEqual(self.client.get(path).status_code, 401)
        self.register()
        self.assertEqual(self.client.post('/api/cvs/draft').status_code, 409)

    def test_draft_manual_versions_and_snapshot(self):
        source = self.profile()
        first = self.draft()
        self.assertEqual(first['origin'], 'profile')
        self.assertEqual(first['document']['full_name'], source['full_name'])
        edited = dict(first['document'], summary='Tóm tắt được chỉnh sửa')
        response = self.client.post(f"/api/cvs/{first['id']}/versions", json={'document': edited})
        self.assertEqual(response.status_code, 201, response.text)
        second = response.json()
        self.assertEqual(second['parent_id'], first['id'])
        self.assertEqual(second['origin'], 'manual')
        self.assertEqual(self.client.get(f"/api/cvs/{first['id']}").json()['document']['summary'], '')
        self.client.put('/api/profile', json=dict(source, full_name='Hồ sơ đổi tên'))
        self.assertEqual(self.client.get(f"/api/cvs/{second['id']}").json()['document']['full_name'], source['full_name'])
        with Session(self.connection) as db:
            self.assertEqual(db.get(CV, second['id']).profile_snapshot['full_name'], source['full_name'])
        self.client.post('/api/auth/logout')
        self.client.post('/api/auth/login', json=self.credentials)
        self.assertEqual(self.client.get('/api/cvs').json()['items'][0]['id'], second['id'])

    def test_cross_account_access(self):
        self.profile()
        version = self.draft()
        self.register(self.other, dict(self.credentials, email='other-cv@example.com'))
        self.assertEqual(self.other.get('/api/cvs').json()['items'], [])
        base = f"/api/cvs/{version['id']}"
        for path in [base, base + '/export/pdf', base + '/export/docx']:
            self.assertEqual(self.other.get(path).status_code, 404)
        self.assertEqual(self.other.post(base + '/versions', json={'document': version['document']}).status_code, 404)
        self.assertEqual(self.other.post(base + '/revise', json={'instructions': 'Viết gọn', 'consent': True}).status_code, 404)

    def test_validation_and_csrf(self):
        self.profile()
        version = self.draft()
        self.assertEqual(self.client.post('/api/cvs/generate', json={'consent': False}).status_code, 422)
        self.assertEqual(self.client.post('/api/cvs/generate', json={'instructions': 'x' * 2001, 'consent': True}).status_code, 422)
        invalid = dict(version['document'], full_name='   ')
        self.assertEqual(self.client.post(f"/api/cvs/{version['id']}/versions", json={'document': invalid}).status_code, 422)
        invalid = dict(version['document'], skills='bad\x00text')
        self.assertEqual(self.client.post(f"/api/cvs/{version['id']}/versions", json={'document': invalid}).status_code, 422)
        self.assertEqual(self.client.post('/api/cvs/draft', headers={'Origin': 'https://evil.example'}).status_code, 403)
        self.assertEqual(self.client.get(f"/api/cvs/{version['id']}/export/html").status_code, 404)

    def test_missing_key_does_not_create_fake_ai_cv(self):
        self.profile()
        with patch.object(settings, 'gemini_api_key', SecretStr('')):
            self.assertFalse(self.client.get('/api/cvs/status').json()['ai_available'])
            before = self.count()
            response = self.client.post('/api/cvs/generate', json={'consent': True})
            self.assertEqual(response.status_code, 503)
            self.assertEqual(self.count(), before)
            self.draft()

    def test_ai_generation_minimizes_contact_data(self):
        source = self.profile()
        with patch.object(settings, 'gemini_api_key', SecretStr('test-not-a-real-key')), \
                patch('app.gemini._post', return_value=self.ai_response()) as post:
            response = self.client.post('/api/cvs/generate', json={'consent': True, 'instructions': 'Viết ngắn gọn'})
            self.assertEqual(response.status_code, 201, response.text)
            self.assertEqual(response.json()['origin'], 'gemini')
            self.assertEqual(response.json()['document']['phone'], source['phone'])
            payload = post.call_args.kwargs['json']
            prompt = json.loads(payload['contents'][0]['parts'][0]['text'])
            self.assertNotIn('full_name', prompt['career_profile'])
            self.assertNotIn(source['phone'], json.dumps(prompt))
            self.assertNotIn(self.email, json.dumps(prompt))
            self.assertNotIn('test-not-a-real-key', response.text)
            self.assertEqual(payload['generationConfig']['responseMimeType'], 'application/json')
            self.assertIn('responseJsonSchema', payload['generationConfig'])

    def test_ai_revision_uses_original_snapshot(self):
        source = self.profile()
        parent = self.draft()
        self.client.put('/api/profile', json=dict(source, skills='Java'))
        with patch.object(settings, 'gemini_api_key', SecretStr('test-not-a-real-key')), \
                patch('app.gemini._post', return_value=self.ai_response()) as post:
            response = self.client.post(f"/api/cvs/{parent['id']}/revise", json={'consent': True, 'instructions': 'Viết súc tích'})
            self.assertEqual(response.status_code, 201, response.text)
            self.assertEqual(response.json()['parent_id'], parent['id'])
            prompt = json.loads(post.call_args.kwargs['json']['contents'][0]['parts'][0]['text'])
            self.assertEqual(prompt['career_profile']['skills'], 'Python, MySQL')
            self.assertNotIn('email', prompt['previous_cv'])
            self.assertEqual(self.client.get(f"/api/cvs/{parent['id']}").json()['origin'], 'profile')

    def test_ai_provider_failures_do_not_save_versions(self):
        self.profile()
        with patch.object(settings, 'gemini_api_key', SecretStr('test-not-a-real-key')):
            for provider_status, expected in [(429, 429), (401, 503), (403, 503), (404, 503), (503, 503), (500, 502), (400, 502)]:
                with self.subTest(status=provider_status), patch('app.gemini._post', return_value=self.ai_response(status=provider_status)):
                    before = self.count()
                    self.assertEqual(self.client.post('/api/cvs/generate', json={'consent': True}).status_code, expected)
                    self.assertEqual(self.count(), before)
            for response in [self.ai_response(finish='MAX_TOKENS'), self.ai_response(data={}),
                             self.ai_response(data={'summary': 'Missing fields'}),
                             httpx.Response(200, json={'candidates': []})]:
                with patch('app.gemini._post', return_value=response):
                    before = self.count()
                    self.assertEqual(self.client.post('/api/cvs/generate', json={'consent': True}).status_code, 502)
                    self.assertEqual(self.count(), before)
            for error, expected in [(httpx.ReadTimeout('timeout'), 504), (httpx.ConnectError('offline'), 502)]:
                with patch('app.gemini._post', side_effect=error):
                    self.assertEqual(self.client.post('/api/cvs/generate', json={'consent': True}).status_code, expected)

    def test_exports_vietnamese_and_untrusted_text(self):
        self.profile()
        version = self.draft()
        document = dict(version['document'], projects='Dự án <script> & XML kiểm thử')
        version = self.client.post(f"/api/cvs/{version['id']}/versions", json={'document': document}).json()
        base = f"/api/cvs/{version['id']}/export"
        pdf = self.client.get(base + '/pdf')
        self.assertEqual(pdf.status_code, 200, pdf.text if pdf.status_code != 200 else '')
        self.assertEqual(pdf.headers['content-type'], 'application/pdf')
        self.assertIn('attachment', pdf.headers['content-disposition'])
        self.assertEqual(pdf.headers['cache-control'], 'no-store')
        text = '\n'.join(page.extract_text() for page in PdfReader(BytesIO(pdf.content)).pages)
        self.assertIn('Nguyễn Minh Anh', text)
        self.assertIn('Dự án <script> & XML kiểm thử', text)
        docx = self.client.get(base + '/docx')
        self.assertEqual(docx.status_code, 200)
        doc = Document(BytesIO(docx.content))
        paragraphs = [p.text for p in doc.paragraphs]
        paragraphs.extend(p.text for table in doc.tables for row in table.rows for cell in row.cells for p in cell.paragraphs)
        self.assertIn('Nguyễn Minh Anh', paragraphs)
        self.assertIn('Dự án <script> & XML kiểm thử', paragraphs)

    def test_long_cv_paginates_without_losing_end(self):
        self.profile()
        version = self.draft()
        text = '\n'.join(f'Dòng {index}: Phát triển ứng dụng Python và kiểm thử dữ liệu.' for index in range(100))
        document = dict(version['document'], experience=text)
        version = self.client.post(f"/api/cvs/{version['id']}/versions", json={'document': document}).json()
        response = self.client.get(f"/api/cvs/{version['id']}/export/pdf")
        reader = PdfReader(BytesIO(response.content))
        self.assertGreater(len(reader.pages), 1)
        self.assertIn('Dòng 99:', '\n'.join(page.extract_text() for page in reader.pages))

    def test_history_pagination(self):
        self.profile()
        first = self.draft()
        with Session(bind=self.connection, join_transaction_mode='create_savepoint') as db:
            parent = db.get(CV, first['id'])
            for _ in range(31):
                db.add(CV(user_id=parent.user_id, profile_snapshot=parent.profile_snapshot, content=parent.content))
            db.commit()
        page = self.client.get('/api/cvs').json()
        self.assertEqual(len(page['items']), 30)
        tail = self.client.get(f"/api/cvs?before_id={page['next_before_id']}").json()
        self.assertEqual(len(tail['items']), 2)
        self.assertIsNone(tail['next_before_id'])
        self.assertFalse({item['id'] for item in page['items']} & {item['id'] for item in tail['items']})

    def test_pdf_splits_both_columns_and_wraps_long_name(self):
        from app.cv_export import export_pdf
        from app.cv_schemas import CVDocument
        document = CVDocument(
            full_name='Nguyễn Hoàng Minh Anh Trần Phương Thảo',
            email='candidate-with-a-long-email-address@example.com',
            skills='\n'.join(f'Kỹ năng {i}: Python và xử lý dữ liệu.' for i in range(90)),
            experience='\n'.join(f'Công việc {i}: Phát triển hệ thống backend.' for i in range(110)))
        reader = PdfReader(BytesIO(export_pdf(document)))
        text = '\n'.join(page.extract_text() for page in reader.pages)
        self.assertGreater(len(reader.pages), 1)
        self.assertIn('Kỹ năng 89:', text)
        self.assertIn('Công việc 109:', text)
        for i in range(90):
            self.assertEqual(text.count(f'Kỹ năng {i}:'), 1)
        for i in range(110):
            self.assertEqual(text.count(f'Công việc {i}:'), 1)

    def upload(self, color='red'):
        output = BytesIO()
        Image.new('RGB', (200, 300), color).save(output, 'PNG')
        response = self.client.post('/api/cvs/photos', content=output.getvalue(), headers={'Content-Type': 'image/png'})
        self.assertEqual(response.status_code, 201, response.text)
        return response.json()['id']

    def test_photo_mysql_storage_validation_and_access(self):
        self.profile()
        photo_id = self.upload()
        self.assertEqual(self.upload(), photo_id)
        photo_response = self.client.get(f'/api/cvs/photos/{photo_id}')
        self.assertEqual(photo_response.status_code, 200)
        self.assertEqual(photo_response.headers['content-type'], 'image/jpeg')
        self.assertEqual(photo_response.headers['cache-control'], 'no-store')
        image = Image.open(BytesIO(photo_response.content))
        self.assertEqual(image.size, (720, 720))
        self.assertEqual(len(image.getexif()), 0)
        with Session(self.connection) as db:
            self.assertEqual(db.get(CVPhoto, photo_id).data, photo_response.content)
        self.register(self.other, dict(self.credentials, email='photo-other@example.com'))
        self.assertEqual(self.other.get(f'/api/cvs/photos/{photo_id}').status_code, 404)
        self.other.put('/api/profile', json={'full_name': 'Other user'})
        self.assertEqual(self.other.post('/api/cvs/draft', json={'design': {'template_id': 'portrait', 'photo_id': photo_id}}).status_code, 404)
        self.assertEqual(self.client.post('/api/cvs/photos', content=b'<svg>not a photo</svg>').status_code, 422)
        self.assertEqual(self.client.post('/api/cvs/photos', content=b'x' * (3 * 1024 * 1024 + 1)).status_code, 413)
        self.client.post('/api/auth/logout')
        self.assertEqual(self.client.get(f'/api/cvs/photos/{photo_id}').status_code, 401)

    def test_templates_photo_required_and_private_version(self):
        self.profile()
        templates = self.client.get('/api/cvs/templates').json()
        self.assertEqual({item['id'] for item in templates}, {'monochrome', 'lavender', 'editorial', 'gradient'})
        self.assertEqual(self.client.post('/api/cvs/draft', json={'design': {'template_id': 'unknown'}}).status_code, 422)
        self.assertEqual(self.client.post('/api/cvs/draft', json={'design': {'template_id': 'portrait'}}).status_code, 422)
        first_photo, second_photo = self.upload('red'), self.upload('blue')
        response = self.client.post('/api/cvs/draft', json={'design': {'template_id': 'portrait', 'photo_id': first_photo}})
        self.assertEqual(response.status_code, 201, response.text)
        first = response.json()
        second = self.client.post(f"/api/cvs/{first['id']}/versions", json={
            'document': first['document'], 'design': {'template_id': 'portrait', 'photo_id': second_photo}}).json()
        self.assertEqual(self.client.get(f"/api/cvs/{first['id']}").json()['design']['photo_id'], first_photo)
        self.assertEqual(second['design']['photo_id'], second_photo)
        inherited = self.client.post(f"/api/cvs/{first['id']}/versions", json={'document': first['document']}).json()
        self.assertEqual(inherited['design'], first['design'])
        minimal = self.client.post(f"/api/cvs/{first['id']}/versions", json={
            'document': first['document'], 'design': {'template_id': 'minimal', 'photo_id': first_photo}}).json()
        self.assertIsNone(minimal['design']['photo_id'])
        for version in (first, minimal):
            base = f"/api/cvs/{version['id']}/export"
            pdf = self.client.get(base + '/pdf')
            self.assertEqual(pdf.status_code, 200)
            reader = PdfReader(BytesIO(pdf.content))
            self.assertIn('Nguyễn Minh Anh', ' '.join(reader.pages[0].extract_text().split()))
            self.assertEqual(bool(list(reader.pages[0].images)), version['design']['template_id'] == 'portrait')
            docx = self.client.get(base + '/docx')
            self.assertEqual(docx.status_code, 200)
            with ZipFile(BytesIO(docx.content)) as archive:
                self.assertEqual(any(name.startswith('word/media/') for name in archive.namelist()), version['design']['template_id'] == 'portrait')

    def test_reference_templates_and_pdf_preview(self):
        self.profile()
        photo_id = self.upload()
        first = self.draft()
        document = dict(first['document'], languages='Tiếng Anh · Đọc tài liệu',
                        certifications='Chứng chỉ Python', interests='Đọc sách')
        before = self.count()
        for template in ('monochrome', 'lavender', 'editorial', 'gradient'):
            response = self.client.post('/api/cvs/preview', json={
                'document': document, 'design': {'template_id': template, 'photo_id': photo_id}})
            self.assertEqual(response.status_code, 200, response.text if response.status_code != 200 else '')
            reader = PdfReader(BytesIO(response.content))
            text = '\n'.join(page.extract_text() for page in reader.pages)
            for value in ('Nguyễn Minh Anh', 'Chứng chỉ Python', 'Đọc sách'):
                self.assertIn(value, ' '.join(text.split()))
            self.assertTrue(list(reader.pages[0].images))
            self.assertEqual(response.headers['cache-control'], 'no-store')
        self.assertEqual(self.count(), before)
        rendered = self.client.post('/api/cvs/preview?representation=images', json={
            'document': document, 'design': {'template_id':'editorial', 'photo_id':photo_id}})
        self.assertEqual(rendered.status_code, 200)
        image = Image.open(BytesIO(base64.b64decode(rendered.json()['pages'][0]['image'])))
        self.assertGreater(image.width, 700)
        self.assertEqual(self.count(), before)
        for template in ('monochrome', 'editorial'):
            response = self.client.post('/api/cvs/preview', json={
                'document': document, 'design': {'template_id': template}})
            self.assertEqual(response.status_code, 200)
            self.assertFalse(list(PdfReader(BytesIO(response.content)).pages[0].images))

    def test_pdf_preview_rejects_foreign_photo_and_requires_consent_header(self):
        self.profile()
        photo_id = self.upload()
        first = self.draft()
        self.register(self.other, dict(self.credentials, email='preview-other@example.com'))
        response = self.other.post('/api/cvs/preview', json={
            'document': first['document'], 'design': {'template_id': 'gradient', 'photo_id': photo_id}})
        self.assertEqual(response.status_code, 404)
        response = self.client.post('/api/cvs/preview', json={'document': first['document']}, headers={'Origin':'https://evil.example'})
        self.assertEqual(response.status_code, 403)

    def test_reference_layouts_split_long_unbroken_paragraphs(self):
        from app.cv_schemas import CVDocument
        from app.cv_export import export_pdf
        image = BytesIO()
        Image.new('RGB', (100, 100), 'gray').save(image, 'PNG')
        document = CVDocument(full_name='Nguyễn Minh Anh', title='Backend Developer',
            summary=('Kinh nghiệm phát triển sản phẩm. ' * 50) + 'SUMMARY_END',
            skills=('Python và cơ sở dữ liệu. ' * 160) + 'SKILLS_END',
            experience=('Thiết kế ứng dụng và kiểm thử dữ liệu. ' * 160) + 'EXPERIENCE_END')
        for template in ('monochrome', 'lavender', 'editorial', 'gradient'):
            raw = export_pdf(document, template, image.getvalue())
            reader = PdfReader(BytesIO(raw))
            self.assertGreater(len(reader.pages), 1)
            text = ' '.join(' '.join(page.extract_text().split()) for page in reader.pages)
            for marker in ('SUMMARY_END', 'SKILLS_END', 'EXPERIENCE_END'):
                self.assertIn(marker, text)

    def test_lavender_heading_accents_do_not_cross_vietnamese_text(self):
        import pymupdf
        from app.cv_export import export_pdf
        from app.cv_schemas import CVDocument
        photo = BytesIO()
        Image.new('RGB', (100, 100), 'gray').save(photo, 'PNG')
        document = CVDocument(full_name='Nguyễn Minh Anh', summary='Giới thiệu ứng viên.',
            skills='Python', experience='Phát triển ứng dụng.', education='Đại học.',
            projects='Dự án CV.', languages='Tiếng Việt', interests='Đọc sách', certifications='Chứng chỉ Python')
        with pymupdf.open(stream=export_pdf(document, 'lavender', photo.getvalue()), filetype='pdf') as pdf:
            checked = 0
            for page in pdf:
                accents = [drawing['rect'] for drawing in page.get_drawings()
                           if drawing.get('fill') and drawing['rect'].height <= 4.1
                           and abs(drawing['fill'][0] - 116 / 255) < .001
                           and abs(drawing['fill'][1] - 119 / 255) < .001]
                self.assertTrue(accents)
                for block in page.get_text('dict')['blocks']:
                    for line in block.get('lines', []):
                        for span in line['spans']:
                            if span['text'] in ('GIỚI THIỆU', 'KỸ NĂNG', 'KINH NGHIỆM', 'HỌC VẤN', 'DỰ ÁN', 'NGÔN NGỮ', 'SỞ THÍCH', 'CHỨNG CHỈ'):
                                checked += 1
                                rect = pymupdf.Rect(span['bbox'])
                                self.assertFalse(any(rect.intersects(accent) for accent in accents), span['text'])
            self.assertEqual(checked, 8)
