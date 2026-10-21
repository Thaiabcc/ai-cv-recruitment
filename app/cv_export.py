"""In-memory exports: private CVs never become public static files."""
from functools import lru_cache
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont, TTFError
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.config import settings

SECTIONS = (('summary', 'Giới thiệu'), ('education', 'Học vấn'), ('skills', 'Kỹ năng'),
            ('experience', 'Kinh nghiệm làm việc'), ('projects', 'Dự án'),
            ('languages', 'Ngôn ngữ'), ('certifications', 'Chứng chỉ'), ('interests', 'Sở thích'))


def contact_line(document):
    return ' | '.join(value for value in (document.email, document.phone, document.address) if value)


def export_docx(document, template_id='forest', photo=None):
    if template_id in ('portrait', 'lavender', 'gradient') and not photo:
        raise ValueError('Portrait template requires a photo')
    output = BytesIO()
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.8), Cm(1.8)
    section.left_margin, section.right_margin = Cm(2), Cm(2)
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'Arial', Pt(10)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.16
    normal.paragraph_format.widow_control = True
    title = doc.styles['Title']
    title.font.name, title.font.size = 'Arial', Pt(24)
    title.font.bold, title.font.color.rgb = True, RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(5)
    heading = doc.styles['Heading 1']
    heading.font.name, heading.font.size = 'Arial', Pt(12)
    heading.font.color.rgb = RGBColor.from_string('174F3B')
    heading.paragraph_format.space_before = Pt(14)
    heading.paragraph_format.space_after = Pt(6)
    heading.paragraph_format.keep_with_next = True
    doc.add_paragraph(document.full_name, 'Title')
    if document.title:
        run = doc.add_paragraph().add_run(document.title)
        run.bold, run.font.size = True, Pt(12)
    if template_id == 'minimal':
        if contact_line(document):
            doc.add_paragraph(contact_line(document))
        destinations = [(doc, SECTIONS)]
    else:
        table = doc.add_table(rows=1, cols=2)
        table.autofit = False
        table.columns[0].width, table.columns[1].width = Cm(5), Cm(12)
        left, right = table.rows[0].cells
        shade = OxmlElement('w:shd')
        shade.set(qn('w:fill'), 'DFDFE8' if template_id == 'lavender' else 'E6E8EC' if template_id in ('portrait', 'gradient') else 'E7E7E7')
        left._tc.get_or_add_tcPr().append(shade)
        if photo and template_id not in ('forest', 'minimal'):
            left.paragraphs[0].add_run().add_picture(BytesIO(photo), width=Cm(3.2))
        if contact_line(document):
            left.add_paragraph('Liên hệ', 'Heading 1')
            for value in (document.email, document.phone, document.address):
                if value:
                    left.add_paragraph(value)
        destinations = [(left, (('summary', 'Giới thiệu'), ('skills', 'Kỹ năng'), ('languages', 'Ngôn ngữ'), ('certifications', 'Chứng chỉ'), ('interests', 'Sở thích'))),
                        (right, (('experience', 'Kinh nghiệm làm việc'), ('projects', 'Dự án'), ('education', 'Học vấn')))]
    for destination, sections in destinations:
        for key, label in sections:
            text = getattr(document, key)
            if not text:
                continue
            destination.add_paragraph(label, 'Heading 1')
            for line in text.splitlines():
                if line.strip():
                    destination.add_paragraph(line.strip())
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    run = footer.add_run('Trang ')
    run.font.size = Pt(8)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    doc.core_properties.title = f'CV {document.full_name}'
    doc.core_properties.author = document.full_name
    doc.save(output)
    return output.getvalue()


@lru_cache(maxsize=1)
def pdf_fonts():
    pairs = []
    if settings.cv_font_path:
        pairs.append((settings.cv_font_path, settings.cv_font_bold_path or settings.cv_font_path))
    else:
        pairs.extend([
            ('C:/Windows/Fonts/arial.ttf', 'C:/Windows/Fonts/arialbd.ttf'),
            ('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'),
        ])
    for regular, bold in pairs:
        if Path(regular).is_file() and Path(bold).is_file():
            try:
                pdfmetrics.registerFont(TTFont('CVRegular', regular))
                pdfmetrics.registerFont(TTFont('CVBold', bold))
            except TTFError as exc:
                raise ValueError('Invalid Unicode font') from exc
            return 'CVRegular', 'CVBold'
    raise ValueError('Unicode fonts not found')


def export_minimal_pdf(document):
    regular, bold = pdf_fonts()
    output = BytesIO()
    pdf = SimpleDocTemplate(output, pagesize=A4, leftMargin=48, rightMargin=48,
                            topMargin=45, bottomMargin=45, title=f'CV {document.full_name}', author=document.full_name)
    body = ParagraphStyle('MinimalBody', fontName=regular, fontSize=10, leading=15, spaceAfter=6,
                          textColor=colors.HexColor('#374151'), splitLongWords=True)
    name = ParagraphStyle('MinimalName', parent=body, fontName=bold, fontSize=25, leading=31, spaceAfter=8)
    role = ParagraphStyle('MinimalRole', parent=body, fontName=bold, fontSize=12, leading=18)
    heading = ParagraphStyle('MinimalSection', parent=body, fontName=bold, fontSize=12, leading=18,
                             spaceBefore=14, spaceAfter=8, keepWithNext=True)
    story = [Paragraph(escape(document.full_name), name)]
    if document.title:
        story.append(Paragraph(escape(document.title), role))
    if contact_line(document):
        story.append(Paragraph(escape(contact_line(document)), body))
    for field, label in SECTIONS:
        if getattr(document, field):
            story.append(Paragraph(label, heading))
            story.extend(Paragraph(escape(line.strip()), body) for line in getattr(document, field).splitlines() if line.strip())
    def footer(canvas, doc):
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor('#CBD5E1'))
        canvas.line(48, 34, A4[0] - 48, 34)
        canvas.setFont(regular, 8)
        canvas.drawRightString(A4[0] - 48, 22, f'Trang {doc.page}')
        canvas.restoreState()
    pdf.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def export_pdf(document, template_id='forest', photo=None):
    if template_id == 'minimal':
        return export_minimal_pdf(document)
    from app.cv_templates import render_reference_pdf
    return render_reference_pdf(document, template_id, photo, pdf_fonts())
